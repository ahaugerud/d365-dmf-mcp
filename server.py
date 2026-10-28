# /// script
# requires-python = ">=3.10"
# dependencies = ["requests", "mcp<2", "openpyxl"]
# ///
"""MVP MCP server: upload a CSV to a D365 F&O Data Management (DMF) import project.

Auth is the signed-in user's `az login`. Config comes from env vars (set them in the
`env` block of the MCP server definition, see .mcp.json.example):
  DMF_BASE_URL, DMF_DEFINITION_GROUP, DMF_COMPANY, DMF_TEMPLATE_ZIP, DMF_EXPORT_GROUP_NAME
  (optional: DMF_POLL_INTERVAL_SECONDS=5, DMF_POLL_TIMEOUT_SECONDS=300)
Relative paths resolve against this script's directory.
"""
import csv, io, json, os, re, subprocess, time, uuid, zipfile
from datetime import datetime
from pathlib import Path

import requests
from openpyxl import Workbook, load_workbook
from mcp.server.fastmcp import FastMCP

ROOT = Path(__file__).resolve().parent


def _env(key: str, default: str | None = None) -> str:
    v = os.environ.get(key, default)
    if v is None:
        raise SystemExit(f"Missing required env var {key}")
    return v


CFG = {
    "base_url": _env("DMF_BASE_URL"),
    "definition_group": _env("DMF_DEFINITION_GROUP"),
    "company": _env("DMF_COMPANY"),
    "template_zip": _env("DMF_TEMPLATE_ZIP"),
    "export_group_name": _env("DMF_EXPORT_GROUP_NAME"),
    "poll_interval_seconds": float(_env("DMF_POLL_INTERVAL_SECONDS", "5")),
    "poll_timeout_seconds": float(_env("DMF_POLL_TIMEOUT_SECONDS", "300")),
}
BASE = CFG["base_url"].rstrip("/")
ACTIONS = BASE + "/data/DataManagementDefinitionGroups/Microsoft.Dynamics.DataEntities."
DONE = ("Succeeded", "PartiallySucceeded", "Failed", "Canceled")

mcp = FastMCP("d365-dmf")


def _path(p: str) -> Path:
    p = Path(p).expanduser()
    return p if p.is_absolute() else ROOT / p


def _session() -> requests.Session:
    token = subprocess.run(
        ["az", "account", "get-access-token", "--resource", BASE, "--query", "accessToken", "-o", "tsv"],
        capture_output=True, text=True, check=True).stdout.strip()
    s = requests.Session()
    s.headers["Authorization"] = f"Bearer {token}"
    return s


def _call(s, action, body):
    r = s.post(ACTIONS + action, json=body)
    if not r.ok:
        raise RuntimeError(f"{action} failed: {r.status_code} {r.text}")
    return r.json()["value"]


def _read_rows(path: Path) -> list[list[str]]:
    """Read a csv/xlsx file into rows of strings (input type by file extension)."""
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        ws = load_workbook(path, read_only=True, data_only=True).active
        return [["" if c is None else str(c) for c in row] for row in ws.iter_rows(values_only=True)]
    text = path.read_bytes().decode("utf-8-sig")
    return list(csv.reader(io.StringIO(text, newline="")))


def _is_excel(fmt: str) -> bool:
    return "EXCEL" in fmt.upper()


def _encode(rows: list[list[str]], fmt: str, sheet: str) -> bytes:
    if _is_excel(fmt):
        wb = Workbook()
        ws = wb.active
        ws.title = sheet[:31] or "Sheet1"
        for r in rows:
            ws.append(r)
        if rows:  # D365's own Excel files carry a named range over the header row
            from openpyxl.utils import get_column_letter
            from openpyxl.workbook.defined_name import DefinedName
            ref = f"'{ws.title}'!$A$1:${get_column_letter(len(rows[0]))}$1"
            wb.defined_names[ws.title] = DefinedName(ws.title, attr_text=ref)
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()
    buf = io.StringIO(newline="")
    csv.writer(buf, lineterminator="\r\n").writerows(rows)  # DMF needs CRLF
    return buf.getvalue().encode("utf-16" if "UNICODE" in fmt.upper() else "utf-8")  # *-Unicode formats are UTF-16


def _fmt(s, fmt: str) -> str:
    """Resolve a free-text format to the D365 source data format name in use on the environment."""
    r = s.get(BASE + "/data/DataManagementDefinitionGroupDetails", params={"$select": "SourceFormat"})
    r.raise_for_status()
    available = sorted({v["SourceFormat"] for v in r.json()["value"] if v["SourceFormat"]})
    match = next((a for a in available if a.lower() == fmt.strip().lower()), None)
    if not match:
        raise ValueError(f"Unknown format '{fmt}'. In use on this D365 environment: {', '.join(available)}")
    if not (_is_excel(match) or match.upper().startswith("CSV")):
        raise ValueError(f"Format '{match}' is not supported by this server; use a CSV or Excel format")
    return match


def _build_package(src: Path, group: str, zip_path: Path, fmt: str = "CSV"):
    """Reuse manifest/header from an exported package; swap in our data file in the given format."""
    ext = ".xlsx" if _is_excel(fmt) else ".csv"
    with zipfile.ZipFile(_path(CFG["template_zip"])) as zin, zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zout:
        data_name = next(n for n in zin.namelist() if n.lower().endswith(".csv"))
        new_name = data_name[:-4] + ext
        sheet = re.sub(r"\W", "_", data_name[:-4])[:31]  # D365 names the Excel sheet after the entity, with underscores
        for item in zin.infolist():
            if item.filename == data_name:
                zout.writestr(new_name, _encode(_read_rows(src), fmt, sheet))
            elif item.filename == "Manifest.xml":  # UTF-16; point at our import project and format
                xml = zin.read(item.filename).decode("utf-16").replace(CFG["export_group_name"], group)
                xml = xml.replace(f"<InputFilePath>{data_name}<", f"<InputFilePath>{new_name}<")
                xml = re.sub(r"<SourceFormat>[^<]*<", f"<SourceFormat>{fmt}<", xml)
                if _is_excel(fmt):
                    xml = re.sub(r"<ExcelSheetName\s*/>", f"<ExcelSheetName>{sheet}</ExcelSheetName>", xml)
                zout.writestr(item, xml.encode("utf-16"))
            else:
                zout.writestr(item, zin.read(item.filename))


def _import_package(s, zip_path: Path, name: str, group: str, company: str, execute: bool) -> str:
    blob = json.loads(_call(s, "GetAzureWriteUrl", {"uniqueFileName": name}))["BlobUrl"]
    requests.put(blob, data=zip_path.read_bytes(), headers={"x-ms-blob-type": "BlockBlob"}).raise_for_status()
    zip_path.unlink(missing_ok=True)
    exec_id = f"mcp-{uuid.uuid4().hex[:12]}"
    _call(s, "ImportFromPackage", dict(packageUrl=blob, definitionGroupId=group, executionId=exec_id,
                                       execute=execute, overwrite=True, legalEntityId=company))
    return exec_id


@mcp.tool()
def import_file(file_path: str, definition_group: str | None = None, company: str | None = None,
                format: str = "CSV", wait: bool = True, dry_run: bool = False) -> str:
    """Import a data file into D365 through an existing DMF data project, as the signed-in az user.

    The input may be a .csv or .xlsx (first sheet, header row first); it is converted to the
    package `format`, which must match the data project's source data format in D365. Free text,
    case-insensitive: "CSV" (default) and "EXCEL" are the common ones; others in use on the
    environment (e.g. "CSV-Unicode") also work if CSV/Excel based. An unknown value returns the available ones.
    Columns must match the entity of the template package. Defaults for definition_group and
    company come from the server env config.
    With wait=True, polls until the import finishes (or the configured timeout); otherwise
    returns the execution id to check with get_import_status. dry_run only builds the package.
    """
    group = definition_group or CFG["definition_group"]
    company = company or CFG["company"]
    src = _path(file_path)
    if not src.is_file():
        raise FileNotFoundError(src)
    name = f"{group}-{datetime.now():%Y%m%d%H%M%S}"
    zip_path = Path(os.environ.get("TMPDIR", "/tmp")) / f"{name}.zip"
    s = _session()
    fmt = _fmt(s, format)
    _build_package(src, group, zip_path, fmt)
    if dry_run:
        return f"Built {zip_path}"

    exec_id = _import_package(s, zip_path, name, group, company, execute=True)
    if not wait:
        return f"Triggered {exec_id}"
    deadline = time.time() + CFG.get("poll_timeout_seconds", 300)
    while (st := _call(s, "GetExecutionSummaryStatus", {"executionId": exec_id})) not in DONE:
        if time.time() > deadline:
            return f"Triggered {exec_id}; still {st} after timeout, use get_import_status"
        time.sleep(CFG.get("poll_interval_seconds", 5))
    return f"{exec_id}: {st}"


@mcp.tool()
def create_data_project(name: str, description: str = "", company: str | None = None, format: str = "CSV") -> str:
    """Create a new DMF *import* data project for the entity of the configured template package.

    Creates the project, then loads the template package into it without executing, which
    defines the entity and field mapping from the package manifest. Fails if the name exists.
    `format` (free text: "CSV" default, "EXCEL", or another D365 source format name) sets the project's source data format.
    Afterwards use import_file(file_path, definition_group=name, format=format).
    """
    company = company or CFG["company"]
    s = _session()
    fmt = _fmt(s, format)
    r = s.post(BASE + "/data/DataManagementDefinitionGroups", json={
        "Name": name, "OperationType": "Import", "ProjectCategory": "Integration",
        "Description": description, "GenerateDataPackage": "No"})
    if not r.ok:
        raise RuntimeError(f"Create project failed: {r.status_code} {r.text}")
    zip_path = Path(os.environ.get("TMPDIR", "/tmp")) / f"{name}-setup-{datetime.now():%Y%m%d%H%M%S}.zip"
    with zipfile.ZipFile(_path(CFG["template_zip"])) as zin:
        _build_package(Path(zin.extract(next(n for n in zin.namelist() if n.lower().endswith(".csv")),
                                        zip_path.parent / "dmf-tpl")), name, zip_path, fmt)
    _import_package(s, zip_path, zip_path.stem, name, company, execute=False)
    return f"Created data project {name}"


@mcp.tool()
def export_data(output_dir: str, definition_group: str | None = None, company: str | None = None) -> str:
    """Run a DMF *export* data project and save the resulting package contents to output_dir.

    Defaults to the template export project (DMF_EXPORT_GROUP_NAME). The data files are saved in
    the project's own source format (CSV, Excel, ...). Returns the saved file paths.
    """
    group = definition_group or CFG["export_group_name"]
    company = company or CFG["company"]
    s = _session()
    exec_id = f"mcp-{uuid.uuid4().hex[:12]}"
    _call(s, "ExportToPackage", dict(definitionGroupId=group, packageName=exec_id, executionId=exec_id,
                                     reExecute=False, legalEntityId=company))
    deadline = time.time() + CFG["poll_timeout_seconds"]
    while (st := _call(s, "GetExecutionSummaryStatus", {"executionId": exec_id})) not in DONE:
        if time.time() > deadline:
            return f"Export {exec_id} still {st} after timeout, use get_import_status"
        time.sleep(CFG["poll_interval_seconds"])
    if st != "Succeeded":
        return f"Export {exec_id}: {st}"
    r = requests.get(_call(s, "GetExportedPackageUrl", {"executionId": exec_id}))
    r.raise_for_status()
    out = _path(output_dir)
    out.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(io.BytesIO(r.content)) as z:
        files = [n for n in z.namelist() if n not in ("Manifest.xml", "PackageHeader.xml")]
        for n in files:
            (out / Path(n).name).write_bytes(z.read(n))
    return "Saved: " + ", ".join(str(out / Path(n).name) for n in files)


@mcp.tool()
def get_import_status(execution_id: str) -> str:
    """Get the status of a DMF execution (NotRun, Executing, Succeeded, PartiallySucceeded, Failed, Canceled), with error details if it failed."""
    s = _session()
    st = _call(s, "GetExecutionSummaryStatus", {"executionId": execution_id})
    if st in ("Failed", "PartiallySucceeded"):
        st += "\nErrors: " + _call(s, "GetExecutionErrors", {"executionId": execution_id})
    return st


if __name__ == "__main__":
    mcp.run()
