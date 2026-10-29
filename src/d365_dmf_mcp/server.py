"""MCP server: import data files into / export data from D365 F&O Data Management (DMF).

Auth is delegated Entra ID as the signed-in user: `az login` if available, otherwise an
interactive browser login (same mechanism as sql-server-mcp), pinned to DMF_AZURE_TENANT_ID.
Config comes from env vars (set them in the `env` block of the MCP server definition,
see .mcp.json.example):
  DMF_AZURE_TENANT_ID, DMF_BASE_URL
  (optional: DMF_POLL_INTERVAL_SECONDS=5, DMF_POLL_TIMEOUT_SECONDS=300)
Everything else (data project, company, entities, files, format) is passed per tool call.
"""
import csv, io, json, os, re, sys, time, uuid, zipfile
from pathlib import Path
from xml.sax.saxutils import escape

import requests
from azure.core.exceptions import ClientAuthenticationError
from azure.identity import AzureCliCredential, InteractiveBrowserCredential
from mcp.server.fastmcp import FastMCP
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from pydantic import BaseModel


def _env(key: str, default: str | None = None) -> str:
    v = os.environ.get(key, default)
    if v is None:
        raise SystemExit(f"Missing required env var {key}")
    return v


TENANT = _env("DMF_AZURE_TENANT_ID")
BASE = _env("DMF_BASE_URL").rstrip("/")
POLL_INTERVAL = float(_env("DMF_POLL_INTERVAL_SECONDS", "5"))
POLL_TIMEOUT = float(_env("DMF_POLL_TIMEOUT_SECONDS", "300"))
ACTIONS = BASE + "/data/DataManagementDefinitionGroups/Microsoft.Dynamics.DataEntities."
DONE = ("Succeeded", "PartiallySucceeded", "Failed", "Canceled")
NS = "http://schemas.microsoft.com/dynamics/2015/01/DataManagement"

mcp = FastMCP("d365-dmf")


class EntityFile(BaseModel):
    entity: str  # DMF entity name as shown in D365, e.g. "Customer groups"
    file_path: str  # .csv or .xlsx; first row = column (field) names


# ---------------------------------------------------------------- auth / http

_credential = None


def _token() -> str:
    """Token for the D365 environment. Resolution (az-cli, falling back to an interactive
    browser login) is deferred to the first call so the stdio handshake isn't blocked on login."""
    global _credential
    scope = BASE + "/.default"
    if _credential is None:
        cli = AzureCliCredential(tenant_id=TENANT)
        try:
            cli.get_token(scope)
            _credential = cli
        except ClientAuthenticationError:
            sys.stderr.write(f"No active `az login` session for tenant {TENANT!r} - "
                             "falling back to an interactive browser login.\n")
            _credential = InteractiveBrowserCredential(tenant_id=TENANT)
    return _credential.get_token(scope).token


def _session() -> requests.Session:
    s = requests.Session()
    s.headers["Authorization"] = f"Bearer {_token()}"
    return s


def _call(s, action, body):
    r = s.post(ACTIONS + action, json=body)
    if not r.ok:
        raise RuntimeError(f"{action} failed: {r.status_code} {r.text}")
    return r.json()["value"]


def _get(s, path, **params):
    r = s.get(BASE + "/data/" + path, params=params)
    if not r.ok:
        raise RuntimeError(f"GET {path} failed: {r.status_code} {r.text}")
    return r.json()["value"]


def _q(v: str) -> str:
    return v.replace("'", "''")


def _details(s, exec_id: str, st: str) -> str:
    """Per-entity status and error messages for executions that didn't fully succeed."""
    if st not in ("Failed", "PartiallySucceeded"):
        return ""
    entities = _call(s, "GetEntityExecutionSummaryStatusList", {"executionId": exec_id})
    errors = _call(s, "GetExecutionErrors", {"executionId": exec_id})
    return f"\nEntities: {', '.join(entities)}\nErrors: {errors}"


def _wait(s, exec_id: str) -> str:
    deadline = time.time() + POLL_TIMEOUT
    while (st := _call(s, "GetExecutionSummaryStatus", {"executionId": exec_id})) not in DONE:
        if time.time() > deadline:
            return f"{exec_id}: still {st} after timeout, use get_execution_status"
        time.sleep(POLL_INTERVAL)
    return f"{exec_id}: {st}{_details(s, exec_id, st)}"


# ------------------------------------------------------------ files / formats

def _path(p: str) -> Path:
    return Path(p).expanduser()


def _read_rows(path: Path) -> list[list[str]]:
    """Read a csv/xlsx file into rows of strings (input type by file extension)."""
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        ws = load_workbook(path, read_only=True, data_only=True).active
        return [["" if c is None else str(c) for c in row] for row in ws.iter_rows(values_only=True)]
    text = path.read_bytes().decode("utf-8-sig")
    return list(csv.reader(io.StringIO(text, newline="")))


def _is_excel(fmt: str) -> bool:
    return "EXCEL" in fmt.upper()


def _sheet_name(entity: str) -> str:
    return re.sub(r"\W", "_", entity)[:31]  # D365 names the Excel sheet after the entity, with underscores


def _encode(rows: list[list[str]], fmt: str, entity: str) -> bytes:
    if _is_excel(fmt):
        wb = Workbook()
        ws = wb.active
        ws.title = _sheet_name(entity)
        for r in rows:
            ws.append(r)
        if rows:  # D365's own Excel files carry a named range over the header row
            ref = f"'{ws.title}'!$A$1:${get_column_letter(len(rows[0]))}$1"
            wb.defined_names[ws.title] = DefinedName(ws.title, attr_text=ref)
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()
    buf = io.StringIO(newline="")
    csv.writer(buf, lineterminator="\r\n").writerows(rows)  # DMF needs CRLF
    return buf.getvalue().encode("utf-16" if "UNICODE" in fmt.upper() else "utf-8")  # *-Unicode formats are UTF-16


def _fmt(s, fmt: str) -> str:
    """Resolve a free-text format to the D365 source data format name in use on the environment."""
    available = sorted({v["SourceFormat"] for v in _get(s, "DataManagementDefinitionGroupDetails", **{"$select": "SourceFormat"})
                        if v["SourceFormat"]})
    match = next((a for a in available if a.lower() == fmt.strip().lower()), None)
    if not match:
        raise ValueError(f"Unknown format '{fmt}'. In use on this D365 environment: {', '.join(available)}")
    if not (_is_excel(match) or match.upper().startswith("CSV")):
        raise ValueError(f"Format '{match}' is not supported by this server; use a CSV or Excel format")
    return match


# ------------------------------------------------------------------- package

def _resolve_entities(s, entities: list[str]) -> dict[str, dict]:
    """Entity name -> {target, unit, level, seq}: AOT target entity and execution order from D365."""
    targets = {}
    for e in entities:
        rows = _get(s, "DataManagementEntities", **{"$filter": f"EntityName eq '{_q(e)}'", "$select": "EntityName,TargetName"})
        if not rows:
            raise ValueError(f"Unknown DMF entity '{e}' (use the entity name as shown in D365 Data management)")
        targets[e] = rows[0]["TargetName"]
    # "Name-ExecutionUnit-Level-Sequence,..." in dependency order
    seq = _call(s, "GetEntitySequence", {"listOfDataEntities": ",".join(entities)})
    order = {}
    for part in filter(None, seq.split(",")):
        name, unit, level, n = part.rsplit("-", 3)
        order[name] = dict(unit=int(unit), level=int(level), seq=int(n))
    return {e: dict(target=targets[e], **order[e]) for e in entities}


def _manifest(group: str, fmt: str, entities: dict[str, dict], filenames: dict[str, str]) -> str:
    items = ""
    for e, info in entities.items():
        sheet = f"<ExcelSheetName>{escape(_sheet_name(e))}</ExcelSheetName>" if _is_excel(fmt) else "<ExcelSheetName />"
        items += (
            "<DataManagementPackageEntityData><DefaultRefreshType>IncrementalPush</DefaultRefreshType>"
            "<Disable>false</Disable><EntityMapList />"
            f"<EntityName>{escape(e)}</EntityName><EntityTransformList i:nil=\"true\" />{sheet}"
            f"<ExecutionUnit>{info['unit']}</ExecutionUnit><FailExecutionUnitOnError>false</FailExecutionUnitOnError>"
            f"<FailLevelOnError>false</FailLevelOnError><InputFilePath>{escape(filenames[e])}</InputFilePath>"
            f"<LevelInExecutionUnit>{info['level']}</LevelInExecutionUnit><QueryData />"
            "<QueryFilter i:nil=\"true\" /><RunBusinessLogic>true</RunBusinessLogic>"
            "<RunBusinessValidation>true</RunBusinessValidation><SampleFilePath />"
            f"<SequenceInLevel>{info['seq']}</SequenceInLevel><SourceFormat>{escape(fmt)}</SourceFormat>"
            f"<TargetEntity>{escape(info['target'])}</TargetEntity><ValidationStatus>Yes</ValidationStatus>"
            "</DataManagementPackageEntityData>")
    return (f'<DataManagementPackageManifest xmlns:i="http://www.w3.org/2001/XMLSchema-instance" xmlns="{NS}">'
            f"<DefinitionGroupName>{escape(group)}</DefinitionGroupName><Description>{escape(group)}</Description>"
            f"<PackageEntityList>{items}</PackageEntityList><ProjectCategory>0</ProjectCategory>"
            '<RulesData i:nil="true" /></DataManagementPackageManifest>')


def _header(group: str) -> str:
    return (f'<DataManagementPackageHeader xmlns:i="http://www.w3.org/2001/XMLSchema-instance" xmlns="{NS}">'
            f"<Description>{escape(group)}</Description><ManifestType>Microsoft.Dynamics.AX.Framework.Tools."
            "DataManagement.Serialization.DataManagementPackageManifest</ManifestType>"
            "<PackageType>DefinitionGroup</PackageType><PackageVersion>2</PackageVersion></DataManagementPackageHeader>")


def _build_package(group: str, fmt: str, entities: dict[str, dict], files: dict[str, Path]) -> bytes:
    ext = ".xlsx" if _is_excel(fmt) else ".csv"
    filenames = {e: re.sub(r'[\\/:*?"<>|]', "_", e) + ext for e in entities}
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("Manifest.xml", _manifest(group, fmt, entities, filenames).encode("utf-16"))
        z.writestr("PackageHeader.xml", _header(group).encode("utf-16"))
        for e, path in files.items():
            z.writestr(filenames[e], _encode(_read_rows(path), fmt, e))
    return out.getvalue()


def _ensure_project(s, name: str, description: str = "", operation: str = "Import") -> bool:
    """Create the data project if missing. Returns True if it was created."""
    if _get(s, "DataManagementDefinitionGroups", **{"$filter": f"Name eq '{_q(name)}'", "$select": "Name"}):
        return False
    r = s.post(BASE + "/data/DataManagementDefinitionGroups", json={
        "Name": name, "OperationType": operation, "ProjectCategory": "Integration",
        "Description": description or name, "GenerateDataPackage": "No"})
    if not r.ok:
        raise RuntimeError(f"Create project failed: {r.status_code} {r.text}")
    return True


# --------------------------------------------------------------------- tools

@mcp.tool()
def import_data(definition_group: str, company: str, files: list[EntityFile], format: str = "CSV",
                wait: bool = True, dry_run: bool = False) -> str:
    """Import one or more data files into D365 through a DMF data project, as the signed-in user.

    `definition_group` is the data project name (created automatically if it doesn't exist),
    `company` the legal entity id (e.g. USMF). `files` lists {entity, file_path} pairs: `entity`
    is the DMF entity name as shown in D365 (e.g. "Customer groups"; find it with search_entities),
    `file_path` a .csv or .xlsx
    whose first row holds the entity's field names (staging column names, e.g. CUSTOMERGROUPID;
    list them with get_entity_fields).
    Several entities can be imported in one call; D365 decides the order from their dependencies.
    `format` is the file format in the package, free text and case-insensitive: "CSV" (default) and
    "EXCEL" are the common ones; other CSV/Excel-based source formats on the environment
    (e.g. "CSV-Unicode") also work. An unknown value returns the available ones.
    Composite and XML-based entities are not supported.
    With wait=True, polls until the import finishes (or the timeout); otherwise returns the
    execution id to check with get_execution_status. dry_run only validates and builds the package.
    """
    if not files:
        raise ValueError("files must contain at least one {entity, file_path}")
    paths = {f.entity: _path(f.file_path) for f in files}
    if len(paths) != len(files):
        raise ValueError("each entity may appear only once per call")
    for p in paths.values():
        if not p.is_file():
            raise FileNotFoundError(p)
    s = _session()
    fmt = _fmt(s, format)
    entities = _resolve_entities(s, list(paths))
    package = _build_package(definition_group, fmt, entities, paths)
    if dry_run:
        return f"OK: package for {', '.join(paths)} would be {len(package)} bytes ({fmt})"
    _ensure_project(s, definition_group)

    name = f"{definition_group}-{time.strftime('%Y%m%d%H%M%S')}"
    blob = json.loads(_call(s, "GetAzureWriteUrl", {"uniqueFileName": name}))["BlobUrl"]
    requests.put(blob, data=package, headers={"x-ms-blob-type": "BlockBlob"}).raise_for_status()
    exec_id = f"mcp-{uuid.uuid4().hex[:12]}"
    _call(s, "ImportFromPackage", dict(packageUrl=blob, definitionGroupId=definition_group, executionId=exec_id,
                                       execute=True, overwrite=True, legalEntityId=company))
    return _wait(s, exec_id) if wait else f"Triggered {exec_id}"


@mcp.tool()
def create_data_project(name: str, description: str = "") -> str:
    """Create an empty DMF import data project. Optional: import_data creates the project itself
    if it doesn't exist, and the first import defines its entities."""
    if not _ensure_project(_session(), name, description):
        raise RuntimeError(f"Data project '{name}' already exists")
    return f"Created data project {name}"


@mcp.tool()
def export_data(output_dir: str, definition_group: str, company: str) -> str:
    """Run an existing DMF *export* data project and save its data files to output_dir.

    The files are saved in the project's own source format (CSV, Excel, ...). Returns the saved paths.
    """
    s = _session()
    exec_id = f"mcp-{uuid.uuid4().hex[:12]}"
    _call(s, "ExportToPackage", dict(definitionGroupId=definition_group, packageName=exec_id, executionId=exec_id,
                                     reExecute=False, legalEntityId=company))
    result = _wait(s, exec_id)
    if not result.endswith("Succeeded"):
        return result
    r = requests.get(_call(s, "GetExportedPackageUrl", {"executionId": exec_id}))
    r.raise_for_status()
    out = _path(output_dir)
    out.mkdir(parents=True, exist_ok=True)
    saved = []
    with zipfile.ZipFile(io.BytesIO(r.content)) as z:
        for n in z.namelist():
            if n not in ("Manifest.xml", "PackageHeader.xml"):
                (out / Path(n).name).write_bytes(z.read(n))
                saved.append(str(out / Path(n).name))
    return "Saved: " + ", ".join(saved)


@mcp.tool()
def get_execution_status(execution_id: str) -> str:
    """Get the status of a DMF execution (NotRun, Executing, Succeeded, PartiallySucceeded, Failed,
    Canceled), with per-entity status and error messages if it failed."""
    s = _session()
    st = _call(s, "GetExecutionSummaryStatus", {"executionId": execution_id})
    return f"{execution_id}: {st}{_details(s, execution_id, st)}"


@mcp.tool()
def list_data_projects(name_filter: str = "", operation_type: str = "") -> str:
    """List existing DMF data projects (name, Import/Export, description).

    `name_filter` is a case-insensitive substring of the name or description; `operation_type`
    optionally restricts to "Import" or "Export". Use the names with import_data / export_data."""
    rows = _get(_session(), "DataManagementDefinitionGroups", **{"$select": "Name,OperationType,Description"})
    f = name_filter.strip().lower()
    hits = sorted((r for r in rows
                   if (not f or f in r["Name"].lower() or f in (r["Description"] or "").lower())
                   and (not operation_type or r["OperationType"].lower() == operation_type.strip().lower())),
                  key=lambda r: r["Name"].lower())
    if not hits:
        return "No matching data projects"
    lines = [f"{r['Name']} [{r['OperationType']}]" + (f" - {r['Description']}" if r["Description"] and r["Description"] != r["Name"] else "")
             for r in hits[:100]]
    return f"{len(hits)} data project(s)" + (" (showing 100, narrow with name_filter)" if len(hits) > 100 else "") + ":\n" + "\n".join(lines)


_FLAGS = {  # tool argument -> DataManagementDefinitionGroupDetails field
    "run_business_validation": "RunBusinessValidation",
    "run_validate_field": "RunValidateField",
    "run_business_logic": "RunBusinessLogic",
    "skip_staging": "SkipStaging",
    "fail_level_on_error": "FailLevelOnError",
    "fail_execution_unit_on_error": "FailExecutionUnitOnError",
}
_DEFAULTS = {"RunBusinessValidation": "Yes", "RunValidateField": "Yes", "RunBusinessLogic": "Yes",
             "SkipStaging": "No", "FailLevelOnError": "No", "FailExecutionUnitOnError": "No"}


def _entity_row(s, group: str, entity: str) -> dict:
    rows = _get(s, "DataManagementDefinitionGroupDetails", **{
        "$filter": f"DefinitionGroupId eq '{_q(group)}' and EntityName eq '{_q(entity)}'",
        "$select": ",".join(_FLAGS.values())})
    if not rows:
        raise ValueError(f"Entity '{entity}' is not in data project '{group}' (see list_project_entities)")
    return rows[0]


def _flags_text(row: dict) -> str:
    changed = [f"{k}={row[k]}" for k, d in _DEFAULTS.items() if row[k] != d]
    return ", ".join(changed) if changed else "default settings"


@mcp.tool()
def list_project_entities(definition_group: str) -> str:
    """List the entities defined in a DMF data project, with source format, execution order
    (unit/level/sequence) and any non-default processing settings (see configure_project_entity)."""
    rows = _get(_session(), "DataManagementDefinitionGroupDetails", **{
        "$filter": f"DefinitionGroupId eq '{_q(definition_group)}'",
        "$select": "EntityName,SourceFormat,ExecutionUnit,LevelInExecutionUnit,SequenceInLevel," + ",".join(_FLAGS.values())})
    if not rows:
        return f"No entities in '{definition_group}' (project missing or empty; see list_data_projects)"
    rows.sort(key=lambda r: (r["ExecutionUnit"], r["LevelInExecutionUnit"], r["SequenceInLevel"]))
    return "\n".join(f"{r['EntityName']} [{r['SourceFormat']}] unit {r['ExecutionUnit']} level {r['LevelInExecutionUnit']} "
                     f"seq {r['SequenceInLevel']} ({_flags_text(r)})" for r in rows)


@mcp.tool()
def configure_project_entity(definition_group: str, entity: str,
                             run_business_validation: bool | None = None,
                             run_validate_field: bool | None = None,
                             run_business_logic: bool | None = None,
                             skip_staging: bool | None = None,
                             fail_level_on_error: bool | None = None,
                             fail_execution_unit_on_error: bool | None = None) -> str:
    """Show or change the processing settings of one entity in a DMF data project. Only the
    arguments you pass are changed (none = just show the current settings); the entity must already
    be in the project (import_data adds it on first use). Settings persist across later imports.

    D365 defaults (and the right choice for almost every import): run_business_validation=True
    (table validateWrite + event handlers), run_validate_field=True (per-field validateField),
    run_business_logic=True (table insert()/update() logic), skip_staging=False,
    fail_level_on_error=False / fail_execution_unit_on_error=False (keep going, report bad rows as
    PartiallySucceeded). RECOMMENDATIONS: leave the defaults. Only turn validations/business logic
    off for large one-off migration loads of data that already imported cleanly in a test run, and
    tell the user first, since the target then accepts data the UI would reject. skip_staging
    (write straight to the target) is faster but only works for entities that support it and gives
    less error detail. Set the fail_* flags to True when later rows or entities are meaningless if
    one fails (level = same dependency level, execution unit = the whole dependency chain).

    Not controllable here (no API; D365 UI only, Data management > Framework parameters > Entity
    settings > Configure entity execution parameters): Import threshold record count (records per
    task) and Import task count (parallel tasks, max = batch threads per server x batch servers),
    plus set-based processing and change tracking on the entity. For big files, point the user to
    those, and note that imports only parallelise when run as a batch job (ImportFromPackage does)."""
    s = _session()
    row = _entity_row(s, definition_group, entity)
    args = locals()
    patch = {field: ("Yes" if args[arg] else "No") for arg, field in _FLAGS.items() if args[arg] is not None}
    if not patch:
        return f"{entity} in {definition_group}: {_flags_text(row)}"
    r = s.patch(BASE + f"/data/DataManagementDefinitionGroupDetails(DefinitionGroupId='{_q(definition_group)}',"
                       f"EntityName='{_q(entity)}')", json=patch)
    if not r.ok:
        raise RuntimeError(f"Update failed: {r.status_code} {r.text}")
    return f"{entity} in {definition_group}: {_flags_text(row)} -> {_flags_text(_entity_row(s, definition_group, entity))}"


@mcp.tool()
def search_entities(query: str, limit: int = 50) -> str:
    """Search the DMF entity catalogue by case-insensitive substring of the entity name, target
    entity or module (e.g. "customer", "AccountsReceivable"). Returns entity names to use with
    import_data, plus category and module. Composite entities (X composite) are not importable here."""
    rows = _get(_session(), "DataManagementEntities", **{"$select": "EntityName,TargetName,Category,Modules,EntityIsEnabled"})
    q = query.strip().lower()
    hits = sorted((r for r in rows if q in f"{r['EntityName']} {r['TargetName']} {r['Modules']}".lower()),
                  key=lambda r: r["EntityName"].lower())
    if not hits:
        return "No matching entities"
    lines = [f"{r['EntityName']} ({r['Category']}, {r['Modules'] or '-'})" + ("" if r["EntityIsEnabled"] == "Yes" else " [disabled]")
             for r in hits[:limit]]
    return f"{len(hits)} entit(y/ies)" + (f", showing {limit}" if len(hits) > limit else "") + ":\n" + "\n".join(lines)


@mcp.tool()
def get_entity_fields(entity: str) -> str:
    """List the importable columns of a DMF entity: the staging column names to use as the header
    row of the data file, with data type, max length and whether D365 flags the field as required.
    System-managed fields are omitted."""
    rows = _get(_session(), "DataManagementTargetMapEntities", **{
        "$filter": f"Entity eq '{_q(entity)}'",
        "$select": "StagingField,IsRequired,FieldDataTypeValue,FieldLength,IsSystem,isFieldEnabled"})
    if not rows:
        return f"Unknown entity '{entity}' (see search_entities)"
    return "\n".join(f"{r['StagingField']} {r['FieldDataTypeValue']}" + (f"({r['FieldLength']})" if r["FieldLength"] else "")
                     + (" required" if r["IsRequired"] == "Yes" else "")
                     for r in rows if r["IsSystem"] != "Yes" and r["isFieldEnabled"] == "Yes")


def main() -> None:
    mcp.run()


if __name__ == "__main__":
    main()
